"""Registration export (CSV / Excel): one row per participant, one column per
registration-form field, plus transaction columns. The DB session is faked;
these tests check the row/column building and the file encoding."""

import base64
import csv
import datetime as dt
import io
import uuid
from contextlib import contextmanager
from decimal import Decimal
from unittest.mock import MagicMock

import pytest
from common.models import (
    Event,
    EventFormField,
    EventStatus,
    FormFieldType,
    Order,
    OrderItem,
    PaymentStatus,
    Ticket,
    TicketTier,
)


def _setup(monkeypatch):
    event = Event(
        id=uuid.uuid4(), organiser_id=uuid.uuid4(), category_id=uuid.uuid4(), title="Jammu BSF Marathon",
        description="x", event_date=dt.date(2026, 11, 22), event_time=dt.time(5, 0), venue_name="Stadium",
        venue_address="Jammu", city="Jammu", capacity=100, status=EventStatus.LIVE,
    )
    tshirt = EventFormField(id=uuid.uuid4(), event_id=event.id, label="T-shirt size", field_type=FormFieldType.SINGLE_CHOICE, options=["S", "M"], required=True, sort_order=1)
    dob = EventFormField(id=uuid.uuid4(), event_id=event.id, label="Date of birth", field_type=FormFieldType.DOB, required=True, sort_order=0)
    event.form_fields = [tshirt, dob]

    full = TicketTier(id=uuid.uuid4(), event_id=event.id, name="Full Marathon", price=Decimal("600"), quantity_total=10, quantity_sold=1)
    half = TicketTier(id=uuid.uuid4(), event_id=event.id, name="Half Marathon", price=Decimal("400"), quantity_total=10, quantity_sold=1)
    order = Order(
        id=uuid.uuid4(), event_id=event.id, buyer_name="Aditi Rao", buyer_email="aditi@example.com", buyer_phone="+919876543210",
        subtotal=Decimal("1000"), booking_fee=Decimal("0"), total_amount=Decimal("900"), discount_amount=Decimal("100"),
        promo_code="RUN10", payment_status=PaymentStatus.SUCCESS, payment_gateway_ref="pay_ABC123",
        created_at=dt.datetime(2026, 10, 1, 9, 30, tzinfo=dt.UTC),
    )
    order.form_responses = []

    def _item(tier, name, answers):
        oi = OrderItem(id=uuid.uuid4(), order_id=order.id, ticket_tier_id=tier.id, quantity=1, unit_price=tier.price)
        oi.ticket_tier = tier
        oi.tickets = [Ticket(id=uuid.uuid4(), order_item_id=oi.id, qr_code_token="t", attendee_name=name,
                             attendee_answers=answers, approval_status="approved", checked_in=False)]
        return oi

    order.order_items = [
        _item(full, "Brijesh Kumar", [
            {"field_id": str(dob.id), "field_label": "Date of birth", "answer": "1990-04-24"},
            {"field_id": str(tshirt.id), "field_label": "T-shirt size", "answer": "M"},
        ]),
        # An answer to a field that was since deleted still gets a column.
        _item(half, "Meera Rao", [{"field_id": str(uuid.uuid4()), "field_label": "Blood group", "answer": "O+"}]),
    ]

    session = MagicMock()
    session.get.return_value = event
    session.execute.return_value.scalars.return_value.all.return_value = [order]

    @contextmanager
    def _get_session():
        yield session

    monkeypatch.setattr("events_service.get_session", _get_session)
    return event


def test_csv_export_has_form_columns_and_transaction_details(monkeypatch):
    import events_service as svc

    event = _setup(monkeypatch)
    out = svc.export_attendees_impl(str(event.id), None, "csv", None)

    assert out["filename"] == "jammu-bsf-marathon-registrations.csv"
    assert out["row_count"] == 2
    text = base64.b64decode(out["data"]).decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(text)))

    header = list(rows[0].keys())
    # Form fields in form order, then the orphaned answer label.
    assert header[10:13] == ["Date of birth", "T-shirt size", "Blood group"]
    first, second = rows
    assert first["Transaction ID"] == "pay_ABC123"
    assert first["Transaction date"] == "01-10-2026 09:30"
    assert first["Participant name"] == "Brijesh Kumar"
    assert first["Race / category"] == "Full Marathon"
    assert first["T-shirt size"] == "M"
    assert first["Registration amount"] == "600.00"
    # Order discount/total spread by ticket price: 600/1000 and 400/1000.
    assert first["Discount amount"] == "60.00"
    assert first["Final amount paid"] == "540.00"
    assert second["Final amount paid"] == "360.00"
    assert first["Promo code"] == "RUN10"
    assert first["Payment status"] == "success"
    assert second["Blood group"] == "O+"
    assert second["T-shirt size"] == ""


def test_xlsx_export_opens_with_header_row(monkeypatch):
    import events_service as svc
    from openpyxl import load_workbook

    event = _setup(monkeypatch)
    out = svc.export_attendees_impl(str(event.id), None, "xlsx", "all")
    wb = load_workbook(io.BytesIO(base64.b64decode(out["data"])))
    ws = wb.active
    assert ws["A1"].value == "Order ID"
    assert ws.max_row == 3


def test_export_rejects_unknown_format_and_status(monkeypatch):
    import events_service as svc
    from aws_lambda_powertools.event_handler.exceptions import BadRequestError

    event = _setup(monkeypatch)
    with pytest.raises(BadRequestError):
        svc.export_attendees_impl(str(event.id), None, "pdf", None)
    with pytest.raises(BadRequestError):
        svc.export_attendees_impl(str(event.id), None, "csv", "paid")
